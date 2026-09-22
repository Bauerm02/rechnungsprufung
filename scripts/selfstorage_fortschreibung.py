"""Read-only monthly workbook check. No AI, invoice, bank import or mail.

Uses the bundled openpyxl runtime. A source date is not a service period:
only explicitly configured source hashes receive a confirmed month.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re


def number(value):
    if value is None or isinstance(value, bool):
        raise ValueError("Pflichtbetrag fehlt oder ist kein Betrag")
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError("Betrag ist nicht endlich")
    return result


def cents(value):
    return int((number(value) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def calculate(revenue, operating, reserve, flat_cost, threshold, lower, upper):
    revenue, operating, reserve, flat_cost, threshold, lower, upper = map(
        number, (revenue, operating, reserve, flat_cost, threshold, lower, upper)
    )
    if min(revenue, operating, reserve, flat_cost, threshold) < 0:
        raise ValueError("Negative Einnahmen oder Kosten benötigen Prüfung")
    if not (0 <= lower <= 1 and 0 <= upper <= 1):
        raise ValueError("Ungültiger Verteilungsschlüssel")
    basis = revenue - operating - reserve - flat_cost
    if basis < 0:
        raise ValueError("Negative Verteilungsbasis benötigt Prüfung")
    owner = min(basis, threshold) * lower + max(basis - threshold, 0) * upper
    operator = basis - owner
    payout = owner + operating + reserve
    return dict(basis=str(basis), owner=str(owner), operator=str(operator),
                payout=str(payout), payout_cent=cents(payout))


FORMULAS = {
    'A3': '=MS!I108', 'A4': '=A3', 'A8': '=SUM(A4:A7)',
    'B14': '=IF(A8<=A11,A8*B11,A11*B11)',
    'B15': '=IF(A8>A11,(A8-A11)*B12,0)', 'B16': '=SUM(B14:B15)',
    'D14': '=IF(A8<=A11,A8*D11,A11*D11)',
    'D15': '=IF(A8>A11,(A8-A11)*D12,0)', 'D16': '=SUM(D14:D15)',
    'D19': '=A3', 'B20': '=A5', 'B21': '=A6', 'D22': '=A7',
    'B23': '=(A3-A4)*-1', 'B25': '=-B16', 'D25': '=-D16',
    'B28': '=SUM(B20:B27)', 'D28': '=SUM(D19:D27)',
}


def inspect(path: Path, period_by_hash: dict):
    import openpyxl
    original = path.read_bytes()
    digest = hashlib.sha256(original).hexdigest()
    period = period_by_hash.get(digest)
    if period is not None and not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', period):
        raise ValueError('Leistungsmonat muss YYYY-MM sein')
    formulas = openpyxl.load_workbook(BytesIO(original), data_only=False, read_only=True)
    values = openpyxl.load_workbook(BytesIO(original), data_only=True, read_only=True)
    try:
        f, v = formulas['Berechnung Focky'], values['Berechnung Focky']
        if str(formulas['MS']['I108'].value).replace(' ','').upper() != '=SUM(I2:I107)':
            raise ValueError('Umsatz-Summenbereich geändert')
        for cell, label in [('A28', 'Eigentümer'), ('B6', 'Rücklage'), ('B7', 'Strom')]:
            if label.casefold() not in str(v[cell].value).casefold():
                raise ValueError(f'{cell}: Layout/Bezeichnung geändert')
        for cell, expected in FORMULAS.items():
            actual = str(f[cell].value).replace(' ', '').upper()
            if actual != expected.upper():
                raise ValueError(f'{cell}: Formel geändert; fachliche Prüfung erforderlich')
        # Do not silently overlook new deductions (including loans) in payout.
        for col, start, stop, allowed in [('D',19,27,{19,22,25}), ('B',20,27,{20,21,23,25})]:
            for row in range(start,stop+1):
                if row not in allowed and v[f'{col}{row}'].value not in (None, 0):
                    raise ValueError(f'{col}{row}: zusätzliche Position benötigt Prüfung')
        for cell in ['A5','A6','A7']:
            if number(v[cell].value) > 0:
                raise ValueError(f'{cell}: Kosten haben falsches Vorzeichen')
        result = calculate(v['A4'].value, -number(v['A5'].value), -number(v['A6'].value),
                           -number(v['A7'].value), v['A11'].value, v['B11'].value, v['B12'].value)
        if number(v['B11'].value)+number(v['D11'].value)!=1 or number(v['B12'].value)+number(v['D12'].value)!=1:
            raise ValueError('Verteilungsschlüssel ergänzen sich nicht auf 100 Prozent')
        detail = sum(number(row[0] or 0) for row in values['MS'].iter_rows(min_row=2,max_row=107,min_col=9,max_col=9,values_only=True))
        if cents(detail) != cents(v['A4'].value):
            raise ValueError('Einzelmieten weichen von Umsatz ab')
        for cell, key in [('A8','basis'),('B16','owner'),('D16','operator'),('D28','payout')]:
            if cents(v[cell].value) != cents(result[key]):
                raise ValueError(f'{cell}: gespeicherter Betrag stimmt rechnerisch nicht')
        if cents(v['B28'].value) != -result['payout_cent']:
            raise ValueError('Gegenrechnung B28 stimmt nicht')
        return dict(source=str(path),sha256=digest,period=period,
                    status='GERECHNET' if period else 'MONAT_ZU_BESTAETIGEN',
                    netto_approved=False,bank_payment=None,**result)
    finally:
        formulas.close(); values.close()


def run(config_path):
    config=json.loads(Path(config_path).read_text(encoding='utf-8-sig'))
    output=Path(config['output_file'])
    rows=[]
    root=Path(config['source_dir']).resolve()
    if not root.is_dir():
        result=dict(checked_at=datetime.now().astimezone().isoformat(),unit=config['unit'],
                    source_available=False,reports=[],error='Quellordner nicht erreichbar')
        output.parent.mkdir(parents=True,exist_ok=True)
        temp=output.with_suffix('.tmp')
        temp.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        temp.replace(output)
        raise FileNotFoundError(str(root))
    for path in sorted(root.glob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        try:
            if hashlib.sha256(path.read_bytes()).hexdigest() in config.get('ignored_historical_hashes', []):
                continue
            rows.append(inspect(path,config.get('period_by_hash',{})))
        except Exception as exc:
            rows.append(dict(source=str(path),status='PRUEFUNG_ERFORDERLICH',error=str(exc)))
    result=dict(checked_at=datetime.now().astimezone().isoformat(),unit=config['unit'],
                source_available=True,reports=rows)
    output.parent.mkdir(parents=True,exist_ok=True)
    temporary=output.with_suffix('.tmp')
    temporary.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    temporary.replace(output)
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',required=True)
    args=parser.parse_args()
    result=run(args.config)
    print(json.dumps({'reports':len(result['reports']), 'checked_at':result['checked_at']}))
